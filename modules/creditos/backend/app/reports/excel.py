"""Exportación a Excel (openpyxl). Reemplaza los 'Exportar a Excel' del VFP."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill, Alignment

AZUL = "FF14428A"


def _encabezar(ws, columnas):
    ws.append(columnas)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor=AZUL)
        cell.alignment = Alignment(horizontal="center")


ESTADO_TXT = {"A": "Activo", "C": "Cancelado"}


def _num(ws, min_col, max_col, min_row=2):
    for row in ws.iter_rows(min_col=min_col, max_col=max_col, min_row=min_row):
        for cell in row:
            cell.number_format = "#,##0.00"


def listado_creditos_excel(items) -> bytes:
    """Listado de créditos: crédito, cliente, CUIL, línea, capital, saldo, estado."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Créditos"
    _encabezar(ws, ["Crédito", "Cliente", "CUIL", "Línea", "Capital", "Saldo", "Estado"])
    tot_cap = tot_sal = 0.0
    for c in items:
        cap, sal = float(c["capital"]), float(c["saldo"])
        tot_cap += cap; tot_sal += sal
        ws.append([c["credito_id"], c["cliente"], c["cuil"], c["linea"],
                   cap, sal, ESTADO_TXT.get(c["estado"], c["estado"])])
    ws.append([])
    ws.append(["", "", "", "TOTAL", tot_cap, tot_sal, ""])
    for row in ws.iter_rows(min_col=5, max_col=6, min_row=2):
        for cell in row:
            cell.number_format = "#,##0.00"
    for idx, wdt in enumerate([9, 34, 13, 30, 15, 15, 11], start=1):
        ws.column_dimensions[chr(64 + idx)].width = wdt
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def pagos_caja_excel(items) -> bytes:
    """Pagos de créditos en caja (write-only: soporta decenas de miles de filas)."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Pagos en caja")
    for idx, wdt in enumerate([13, 9, 7, 34, 9, 8, 14, 15], start=1):
        ws.column_dimensions[chr(64 + idx)].width = wdt

    hdr = []
    for name in ["Fecha pago", "Crédito", "Cuota", "Cliente", "Recibo",
                 "Vía", "Cajero", "Importe"]:
        c = WriteOnlyCell(ws, value=name)
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center")
        hdr.append(c)
    ws.append(hdr)

    total = 0.0
    for p in items:
        imp = float(p["total_pagado"])
        total += imp
        importe = WriteOnlyCell(ws, value=imp)
        importe.number_format = "#,##0.00"
        ws.append([str(p["fecha_pago"] or ""), p["credito_id"], p["cuota"],
                   p["cliente"], p["nro_recibo"], p["via_pago"], p["cajero"], importe])
    tot_cell = WriteOnlyCell(ws, value=total)
    tot_cell.number_format = "#,##0.00"
    ws.append(["", "", "", "", "", "", "TOTAL", tot_cell])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def turnos_excel(items) -> bytes:
    """Turnos otorgados de crédito (write-only): fecha, período, tipo, N°,
    solicitante, CUIL, sueldo, usado, autorizado."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Turnos")
    for idx, wdt in enumerate([12, 9, 8, 8, 34, 13, 14, 8, 12], start=1):
        ws.column_dimensions[chr(64 + idx)].width = wdt
    hdr = []
    for name in ["Fecha", "Período", "Tipo", "N° turno", "Solicitante", "CUIL",
                 "Sueldo", "Usado", "Autorizado"]:
        c = WriteOnlyCell(ws, value=name)
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center")
        hdr.append(c)
    ws.append(hdr)
    for t in items:
        sueldo = WriteOnlyCell(ws, value=float(t["sueldo"]))
        sueldo.number_format = "#,##0.00"
        ws.append([str(t["fecha"] or ""), t["periodo"], t["tipo"], t["numero"],
                   t["apellido_nombre"], t["cuil"], sueldo,
                   "Sí" if t["usado"] else "No", "Sí" if t["autorizado"] else "No"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def envios_excel(envios) -> bytes:
    """Padrón de débito por planilla (para el banco): CBU, CUIL, cliente, importe."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Débitos"
    _encabezar(ws, ["CBU", "CUIL", "Cliente", "Crédito", "Cuota", "Vencimiento", "Importe"])
    total = 0
    for i in envios["items"]:
        imp = float(i["importe"])
        total += imp
        ws.append([i["cbu"], i["cuil"], i["cliente"], i["credito_id"],
                   i["cuota_numero"], str(i["vencimiento"]), imp])
    ws.append([])
    ws.append(["", "", "", "", "", "TOTAL", total])
    for row in ws.iter_rows(min_col=7, max_col=7, min_row=2):
        for c in row:
            c.number_format = "#,##0.00"
    widths = [26, 13, 34, 9, 7, 13, 14]
    for idx, wdt in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = wdt

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def contratos_pp_excel(items) -> bytes:
    """Cartera de contratos de Configurar Créditos: nº, cliente, línea, monto, saldo, tasa, estado."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Cartera"
    _encabezar(ws, ["Contrato", "Cliente", "Línea", "Sistema", "Monto", "Saldo capital",
                    "TNA %", "Plazo", "Estado", "Fecha valor"])
    tot_monto = tot_saldo = 0.0
    for c in items:
        monto, saldo = float(c["monto_original"]), float(c["saldo_capital"])
        tot_monto += monto; tot_saldo += saldo
        snap = c.get("snapshot") or {}
        ws.append([c["numero_contrato"], c["cliente_nombre"], snap.get("codigo", ""), c["sistema"],
                   monto, saldo, float(c["tasa"]), c["plazo"], c["estado"], c["fecha_valor"]])
    ws.append([])
    ws.append(["", "", "", "TOTAL", tot_monto, tot_saldo, "", "", "", ""])
    for row in ws.iter_rows(min_col=5, max_col=6, min_row=2):
        for cell in row:
            cell.number_format = "#,##0.00"
    for idx, wdt in enumerate([16, 30, 12, 11, 15, 15, 9, 8, 12, 12], start=1):
        ws.column_dimensions[chr(64 + idx)].width = wdt
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
